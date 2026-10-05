"""
Автоматические проверки расчётов. Запуск из папки проекта одной командой:

    python -m unittest -v

Каждая проверка сравнивает результат программы с ответом, который известен заранее
(посчитан вручную или следует из определения).
"""
import io
import sys
import unittest
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).parent.parent))
import app as web          # noqa: E402
import build_dataset       # noqa: E402
import portfolio as pf     # noqa: E402
import reader             # noqa: E402

NINE = tuple(a for a in pf.IDS if a != "M2")


class TestData(unittest.TestCase):
    def test_table_has_ten_indicators_and_cpi(self):
        df = pf.load_data()
        self.assertEqual(list(df.columns), pf.IDS + ["CPI"])
        self.assertEqual(len(pf.IDS), 10)
        self.assertEqual(pf.IDS[0], "M2")               # индикатор №1 — денежная масса
        self.assertAlmostEqual(df["CPI"].iloc[0], 1.0)  # индекс цен начинается с 1

    def test_real_return_formula(self):
        """Реальная доходность = (1 + рублёвая) / (1 + инфляция) - 1, проверяем на одном месяце."""
        df = pf.load_data()
        nominal = df["MCFTR"]["2019-02"] / df["MCFTR"]["2019-01"] - 1
        inflation = df["CPI"]["2019-02"] / df["CPI"]["2019-01"] - 1
        expected = (1 + nominal) / (1 + inflation) - 1
        real = pf.real_returns(["MCFTR", "GOLD"])
        self.assertAlmostEqual(real["MCFTR"]["2019-02"], expected, places=12)

    def test_common_period(self):
        """Берутся только месяцы, где данные есть по всем выбранным индикаторам."""
        self.assertEqual(pf.real_returns(NINE).index[0], "2017-01")              # ограничивает недвижимость
        self.assertEqual(pf.real_returns(["GOLD", "REALTY"]).index[0], "2017-01")
        self.assertEqual(pf.real_returns(["GOLD", "USD"]).index[0], "2013-01")
        self.assertFalse(pf.real_returns(NINE).isna().any().any())

    def test_cpi_matches_official_annual_inflation(self):
        """Годовая инфляция по нашему индексу цен совпадает с официальной (Банк России) в пределах 0,1 п.п."""
        cpi = pf.load_data()["CPI"]
        official = pd.read_csv(build_dataset.RAW / "cbr_cpi_yoy.csv", index_col=0)["cpi_yoy_pct"]
        ours = (cpi / cpi.shift(12) - 1) * 100
        diff = (ours - official).dropna()
        self.assertGreater(len(diff), 100)
        self.assertLess(diff.abs().max(), 0.1)

    def test_income_adds_to_return(self):
        """Аренда 6% в год добавляет к недвижимости ровно (1 + 0,06/12) каждый месяц."""
        base = pf.prepare(("GOLD", "REALTY"))[1]
        rent = pf.prepare(("GOLD", "REALTY"), None, None, 0.06, 0.0)[1]
        ratio = (1 + rent["REALTY"]) / (1 + base["REALTY"])
        self.assertTrue(np.allclose(ratio, 1 + 0.06 / 12))
        self.assertTrue(np.allclose(rent["GOLD"], base["GOLD"]))     # золото дохода не приносит

    def test_moex_file_reader(self):
        """Файлы, скачанные с сайта биржи, читаются: начальное значение MESMTR — 896,03 на конец 2013 года."""
        series = build_dataset.moex_index("MESMTR")
        self.assertEqual(str(series.index[0]), "2013-12")
        self.assertAlmostEqual(series.iloc[0], 896.03)
        self.assertAlmostEqual(pf.load_data().loc["2013-12", "MESMTR"], 896.03)
        cpi = build_dataset.build_cpi()                              # файл Росстата: январь 2013 г. — 100,97% к декабрю
        self.assertAlmostEqual(cpi.iloc[1], 1.0097)


class TestOptimization(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.real, cls.nominal = pf.prepare(NINE)
        cls.mu, cls.cov = pf.annual_stats(cls.real)
        rng = np.random.default_rng(0)                               # 30 000 случайных портфелей для сравнения
        cls.random = np.vstack([rng.dirichlet(np.full(9, a), 10000) for a in (0.1, 0.3, 1.0)])
        cls.random_risk = np.sqrt(np.einsum("ij,jk,ik->i", cls.random, cls.cov, cls.random))
        cls.random_ret = cls.random @ cls.mu

    def check_weights(self, w):
        self.assertAlmostEqual(w.sum(), 1.0, places=9)               # сумма долей = 100%
        self.assertTrue((w >= 0).all() and (w <= 1).all())           # без отрицательных долей

    def test_two_assets_have_exact_answer(self):
        """Для двух несвязанных активов минимум риска известен из формулы: w1 = s2² / (s1² + s2²)."""
        cov = np.diag([0.04, 0.01])
        w, _ = pf.optimize(np.array([0.1, 0.05]), cov, "min_vol")
        self.assertAlmostEqual(w[0], 0.2, places=5)
        self.assertAlmostEqual(pf.portfolio_risk(w, cov), np.sqrt(0.2 ** 2 * 0.04 + 0.8 ** 2 * 0.01), places=7)

    def test_min_vol_beats_random_portfolios(self):
        w, _ = pf.optimize(self.mu, self.cov, "min_vol")
        self.check_weights(w)
        self.assertLessEqual(pf.portfolio_risk(w, self.cov), self.random_risk.min() + 1e-9)

    def test_min_vol_is_below_every_single_asset(self):
        """Диверсификация: риск портфеля не выше риска самого спокойного актива."""
        w, _ = pf.optimize(self.mu, self.cov, "min_vol")
        self.assertLessEqual(pf.portfolio_risk(w, self.cov), np.sqrt(np.diag(self.cov)).min())
        # Пояснения: вклады в сумме дают доходность и риск портфеля; у вошедших индикаторов бета = 1,
        # у не вошедших — не меньше 1 (их добавление увеличило бы риск)
        e = pf.explain_portfolio(w, self.mu, self.cov)
        self.assertAlmostEqual(sum(e["ret"]), pf.portfolio_return(w, self.mu))
        self.assertAlmostEqual(sum(e["risk"]), pf.portfolio_risk(w, self.cov))
        for share, beta in zip(w, e["beta"]):
            self.assertGreaterEqual(beta, 1 - 1e-3)
            if share > 0.01:
                self.assertAlmostEqual(beta, 1, places=2)

    def test_max_sharpe_beats_random_portfolios(self):
        rf = 0.05
        w, _ = pf.optimize(self.mu, self.cov, "max_sharpe", rf)
        self.check_weights(w)
        best = (pf.portfolio_return(w, self.mu) - rf) / pf.portfolio_risk(w, self.cov)
        self.assertGreaterEqual(best, ((self.random_ret - rf) / self.random_risk).max() - 1e-9)

    def test_target_return_is_met(self):
        w, note = pf.optimize(self.mu, self.cov, "target_return", 0.05, 0.08)
        self.check_weights(w)
        self.assertEqual(note, "")
        self.assertAlmostEqual(pf.portfolio_return(w, self.mu), 0.08, places=6)
        near = np.abs(self.random_ret - 0.08) < 0.0005               # случайные портфели с той же доходностью
        self.assertLessEqual(pf.portfolio_risk(w, self.cov), self.random_risk[near].min() + 2e-4)

    def test_target_risk_is_met(self):
        w, note = pf.optimize(self.mu, self.cov, "target_risk", 0.05, 0.10)
        self.check_weights(w)
        self.assertLessEqual(pf.portfolio_risk(w, self.cov), 0.10 + 1e-6)
        allowed = self.random_risk <= 0.10
        self.assertGreaterEqual(pf.portfolio_return(w, self.mu), self.random_ret[allowed].max() - 1e-9)

    def test_impossible_targets_give_warning(self):
        _, note = pf.optimize(self.mu, self.cov, "target_return", 0.05, 0.90)
        self.assertIn("недостижима", note)
        w, note = pf.optimize(self.mu, self.cov, "target_risk", 0.05, 0.001)
        self.assertIn("ниже минимально возможного", note)
        w_min, _ = pf.optimize(self.mu, self.cov, "min_vol")
        self.assertTrue(np.allclose(w, w_min))

    def test_frontier_goes_up_and_right(self):
        frontier = pf.frontier_for(NINE)
        self.assertEqual(len(frontier), 30)
        risks = [p[0] for p in frontier]
        returns = [p[1] for p in frontier]
        self.assertTrue(all(b >= a - 1e-9 for a, b in zip(risks, risks[1:])))
        self.assertTrue(all(b >= a - 1e-9 for a, b in zip(returns, returns[1:])))
        self.assertAlmostEqual(returns[-1], self.mu.max(), places=6)  # правый конец — самый доходный актив

    def test_frontier_is_cached(self):
        self.assertIs(pf.frontier_for(NINE), pf.frontier_for(NINE))   # второй раз возвращается тот же объект


class TestCurveMetrics(unittest.TestCase):
    """Показатели кривой проверяем на придуманном ряду, где ответ легко посчитать вручную."""

    def setUp(self):
        # один актив: +10%, -50%, +20%, +100%  ->  кривая 1; 1,1; 0,55; 0,66; 1,32
        self.r = pd.DataFrame({"A": [0.10, -0.50, 0.20, 1.00]})
        self.w = np.array([1.0])

    def test_curve_drawdown_recovery(self):
        curve, m = pf.curve_and_metrics(self.r, self.w)
        self.assertTrue(np.allclose(curve, [1, 1.1, 0.55, 0.66, 1.32]))
        self.assertAlmostEqual(m["max_drawdown"], -0.5)               # с 1,10 до 0,55
        self.assertEqual(m["max_recovery_months"], 2)                 # два месяца ниже вершины 1,10
        self.assertTrue(m["recovered"])
        self.assertAlmostEqual(m["mean_return"], 0.2 * 12)
        self.assertAlmostEqual(m["final_value"], 1.32)
        # Пояснения к тем же показателям: где вершина, где дно, когда был самый долгий провал
        e = pf.explain_curve(curve, ["m0", "m1", "m2", "m3", "m4"], self.r, self.r, self.w)
        self.assertEqual(e["dd_top"][0], "m1")
        self.assertEqual(e["dd_trough"][0], "m2")
        self.assertEqual(e["recovery"], ["m1", "m3"])
        self.assertEqual(e["loss_months"], 1)
        self.assertAlmostEqual(e["monthly_mean"] * 12, m["mean_return"])
        self.assertAlmostEqual(e["final_real"], 1.32)

    def test_not_recovered(self):
        _, m = pf.curve_and_metrics(pd.DataFrame({"A": [0.10, -0.50, 0.20]}), self.w)
        self.assertFalse(m["recovered"])
        self.assertEqual(m["max_recovery_months"], 2)

    def test_downside_counts_only_losses(self):
        _, m = pf.curve_and_metrics(self.r, self.w)
        self.assertAlmostEqual(m["downside"], np.sqrt(0.5 ** 2 / 4) * np.sqrt(12))
        _, up = pf.curve_and_metrics(pd.DataFrame({"A": [0.1, 0.2, 0.3]}), self.w)
        self.assertEqual(up["downside"], 0.0)                         # без падений риск потерь равен нулю
        self.assertIsNone(up["sortino"])

    def test_fee_and_tax(self):
        """Один актив: возвращать доли не нужно, комиссия берётся только при покупке."""
        _, m = pf.curve_and_metrics(self.r, self.w, self.r, fee=0.01, tax=0.13)
        self.assertAlmostEqual(m["final_value"], 1.32 * 0.99)
        gain = 1.32 * 0.99 - 1
        self.assertAlmostEqual(m["final_after_tax"], 1.32 * 0.99 - 0.13 * gain)

    def test_fee_reduces_result_of_real_portfolio(self):
        real, nominal = pf.prepare(NINE)
        w = np.full(9, 1 / 9)
        _, free = pf.curve_and_metrics(real, w, nominal)
        _, paid = pf.curve_and_metrics(real, w, nominal, fee=0.001)
        self.assertLess(paid["final_value"], free["final_value"])
        self.assertGreater(paid["final_value"], free["final_value"] * 0.97)   # комиссия 0,1% не съедает больше 3%

    def test_no_tax_on_loss(self):
        loss = pd.DataFrame({"A": [-0.2, 0.1]})
        _, m = pf.curve_and_metrics(loss, self.w, loss, tax=0.13)
        self.assertAlmostEqual(m["final_after_tax"], m["final_value"])


class TestStability(unittest.TestCase):
    def test_windows(self):
        windows = pf.stability(NINE, "min_vol")
        self.assertEqual(len(windows), 10)
        self.assertEqual(windows[-1]["to"], "2026-08")
        for w in windows:
            self.assertAlmostEqual(sum(w["weights"]), 1.0, places=3)

    def test_short_period_is_rejected(self):
        with self.assertRaises(ValueError):
            pf.stability(("GOLD", "REALTY"), "min_vol", start="2024-01")


class TestWebsite(unittest.TestCase):
    """Проверяем сервер так, как к нему обращается страница."""

    def setUp(self):
        self.client = web.app.test_client()
        self.body = {"assets": list(NINE), "criterion": "min_vol", "rf": 0.05}

    def test_pages_open(self):
        self.assertEqual(self.client.get("/").status_code, 200)
        meta = self.client.get("/api/meta").get_json()
        self.assertEqual(len(meta["assets"]), 10)

    def test_optimize(self):
        r = self.client.post("/api/optimize", json=self.body).get_json()
        self.assertAlmostEqual(sum(r["weights"].values()), 1.0, places=3)
        self.assertEqual(r["period"], ["2017-01", "2026-08", 116])
        self.assertEqual(len(r["curve"]["values"]), 117)
        self.assertEqual(len(r["frontier"]), 30)
        t = r["trace"]                                               # пояснения к расчёту
        self.assertEqual((t["period"]["kept"], t["period"]["limiting"]), (116, "REALTY"))
        self.assertEqual(t["period"]["asked"] - t["period"]["dropped"], 116)
        self.assertAlmostEqual(t["curve"]["final_real"], r["metrics"]["final_value"])   # без комиссий совпадают
        self.assertAlmostEqual(sum(t["portfolio"]["ret"]), r["metrics"]["mean_return"])
        self.assertEqual(len(t["correlation"]), 9)
        m2 = r["benchmark"]                                          # ориентир М2: кривая той же длины, начинается с 1
        self.assertEqual((len(m2["values"]), m2["values"][0]), (117, 1.0))

    def test_errors_are_readable(self):
        r = self.client.post("/api/optimize", json={"assets": ["GOLD"]})
        self.assertEqual(r.status_code, 400)
        self.assertIn("два индикатора", r.get_json()["error"])
        r = self.client.post("/api/optimize", json={**self.body, "criterion": "target_risk"})
        self.assertEqual(r.status_code, 400)

    def test_stability_and_export(self):
        r = self.client.post("/api/stability", json=self.body).get_json()
        self.assertEqual(len(r["windows"]), 10)
        f = self.client.post("/api/future", json=self.body).get_json()      # проверка на «будущем»
        self.assertEqual((f["train"], f["test"]), (["2017-01", "2023-12", 84], ["2024-01", "2026-08", 32]))
        self.assertAlmostEqual(sum(f["weights"]), 1.0, places=3)
        self.assertEqual(len(f["portfolio"]["curve"]), 33)
        self.assertAlmostEqual(f["equal"]["return"], pf.real_returns(NINE).loc["2024-01":].mean().mean() * 12)   # равные доли — среднее по индикаторам
        self.assertEqual(self.client.post("/api/future", json={**self.body, "split": "2017-06"}).status_code, 400)  # слишком мало месяцев
        r = self.client.post("/api/export", json=self.body)
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(r.data))
        self.assertEqual(book.sheetnames, ["Портфель", "Кривая доходности", "Эффективная граница", "Индикаторы"])

    def test_upload_gives_same_result_as_builtin_data(self):
        """Скачиваем пример файла с сайта, загружаем его обратно и получаем тот же портфель."""
        example = self.client.get("/api/template").data
        up = self.client.post("/api/upload", data={"files": (io.BytesIO(example), "primer.csv")}).get_json()
        self.assertEqual(len(up["assets"]), 10)
        self.assertEqual(up["report"]["inflation"]["mode"], "file")        # столбец с индексом цен найден сам
        ids = [a["id"] for a in up["assets"][1:]]                           # те же девять, без М2
        mine = self.client.post("/api/optimize", json={**self.body, "dataset": up["dataset"], "assets": ids}).get_json()
        base = self.client.post("/api/optimize", json=self.body).get_json()
        self.assertEqual(mine["period"], base["period"])
        self.assertAlmostEqual(mine["metrics"]["mean_return"], base["metrics"]["mean_return"], places=5)
        self.assertAlmostEqual(mine["metrics"]["volatility"], base["metrics"]["volatility"], places=5)

    def test_forgotten_table_can_be_restored(self):
        example = self.client.get("/api/template").data
        up = self.client.post("/api/upload", data={"files": (io.BytesIO(example), "primer.csv")}).get_json()
        body = {**self.body, "dataset": up["dataset"], "assets": ["C2", "C6"]}
        pf.DATASETS.clear()                                                 # «сервер перезапустился»
        self.assertEqual(self.client.post("/api/optimize", json=body).status_code, 410)
        back = self.client.post("/api/restore", json={k: up[k] for k in ("table", "assets", "report")}).get_json()
        self.assertEqual(back["dataset"], up["dataset"])                    # та же таблица — тот же номер
        self.assertEqual(self.client.post("/api/optimize", json=body).status_code, 200)

    def test_bad_file_gives_readable_error(self):
        r = self.client.post("/api/upload", data={"files": (io.BytesIO("привет\nэто не таблица".encode()), "x.csv")})
        self.assertEqual(r.status_code, 400)
        self.assertIn("не найдена таблица", r.get_json()["error"])


class TestReader(unittest.TestCase):
    """Один и тот же набор данных записываем в файлы разного устройства. Что бы ни пришло, прочитаться должно одно и то же."""

    @classmethod
    def setUpClass(cls):
        df = pf.builtin_data()
        cls.cpi = df["CPI"]
        cls.levels = df[["MCFTR", "GOLD", "USD"]]
        cls.dates = pd.to_datetime(cls.levels.index) + pd.offsets.MonthEnd(0)       # последний день каждого месяца

    def read(self, *files, **options):
        return reader.read_files(list(files), builtin_cpi=self.cpi, **options)

    def check(self, table, column, source, rtol=1e-4):
        """Столбец прочитанной таблицы повторяет встроенный ряд (с точностью до масштаба)."""
        got, want = table[column].dropna(), self.levels[source].reindex(table[column].dropna().index)
        self.assertTrue(np.allclose(got / got.iloc[0], want / want.iloc[0], rtol=rtol))

    def test_semicolon_decimal_comma_russian_dates(self):
        text = "Дата;Акции;Золото;Доллар\n" + "\n".join(
            f"{d:%d.%m.%Y};" + ";".join(f"{x:.2f}".replace(".", ",") for x in row) for d, row in zip(self.dates, self.levels.values))
        table, assets, report = self.read(("a.csv", text.encode("cp1251")))          # кодировка Windows
        self.assertEqual([a["name"] for a in assets], ["Акции", "Золото", "Доллар"])
        self.assertEqual(list(table.index), list(self.levels.index))
        self.check(table, "C1", "MCFTR")
        self.assertIn("Windows-1251", report["files"][0]["format"])
        self.assertEqual(report["inflation"]["mode"], "builtin")                     # своей инфляции нет — взята встроенная

    def test_comma_separated_decimal_dot_iso_and_us_dates(self):
        for pattern in ("%Y-%m-%d", "%m/%d/%Y"):
            text = "date,stocks,gold,usd\n" + "\n".join(
                f"{d.strftime(pattern)}," + ",".join(f"{x:.4f}" for x in row) for d, row in zip(self.dates, self.levels.values))
            table, _, _ = self.read(("b.csv", text.encode()))
            self.assertEqual(list(table.index), list(self.levels.index))
            self.check(table, "C2", "GOLD")

    def test_title_lines_thousands_separators_and_utf16(self):
        text = "Отчёт по индексам\nИсточник: тест\n\nМесяц\tАкции\tЗолото\n" + "\n".join(
            f"{d:%m.%Y}\t" + f"{row[0]:,.2f}\t{row[1]:,.2f}".replace(",", " ").replace(".", ",") for d, row in zip(self.dates, self.levels.values))
        table, assets, _ = self.read(("c.txt", text.encode("utf-16")))
        self.assertEqual(len(assets), 2)
        self.check(table, "C1", "MCFTR")                                             # «1 718,76» прочитано как 1718,76

    def test_long_table_is_pivoted(self):
        rows = ["TRADEDATE;SECID;CLOSE;VOLUME"] + [f"{d:%Y-%m-%d};{name};{value:.2f};1000"
                for d, row in zip(self.dates, self.levels.values) for name, value in zip(self.levels.columns, row)]
        table, assets, report = self.read(("d.csv", "\n".join(rows).encode()))
        self.assertEqual(sorted(a["name"] for a in assets), ["GOLD", "MCFTR", "USD"])
        self.assertIn("длинная", report["files"][0]["tables"][0]["layout"])

    def test_single_asset_files_are_combined(self):
        """Выгрузка свечей с биржи и файл «прошлые данные» с Investing: из каждого берётся цена закрытия."""
        moex = "\n".join(["candles", "", "open;close;high;low;value;volume;begin;end"] + [
            f"{v * .98:.2f};{v:.2f};{v * 1.02:.2f};{v * .95:.2f};0;0;{d:%Y-%m}-01 00:00:00;{d:%Y-%m-%d} 23:59:59"
            for d, v in zip(self.dates, self.levels["MCFTR"])])
        money = lambda v: f"{v:,.2f}".replace(",", " ").replace(".", ",").replace(" ", ".")       # 12.574,30
        investing = '"Дата","Цена","Откр.","Макс.","Мин.","Объём","Изм. %"\n' + "\n".join(
            f'"{d:%d.%m.%Y}","{money(v)}","{money(v)}","{money(v)}","{money(v)}","1,2M","0,5%"'
            for d, v in list(zip(self.dates, self.levels["GOLD"]))[::-1])                           # новые строки сверху
        table, assets, _ = self.read(("MCFTR.csv", moex.encode()), ("Золото.csv", investing.encode()))
        self.assertEqual([a["name"] for a in assets], ["MCFTR", "Золото"])           # названия взяты из имён файлов
        self.check(table, "C1", "MCFTR")
        self.check(table, "C2", "GOLD")
        with self.assertRaises(reader.ReadError):                                    # один индикатор — портфеля не будет
            self.read(("MCFTR.csv", moex.encode()))

    def test_excel_with_extra_sheet_offset_and_daily_data(self):
        days = pd.date_range("2018-01-01", "2023-12-31", freq="B")
        rng = np.random.default_rng(1)
        prices = pd.DataFrame({"Дата": days, "Фонд А": 100 * np.cumprod(1 + rng.normal(0.0004, 0.01, len(days))),
                               "Фонд Б": 50 * np.cumprod(1 + rng.normal(0.0002, 0.004, len(days)))})
        book = io.BytesIO()
        with pd.ExcelWriter(book, engine="openpyxl") as writer:
            pd.DataFrame({"x": ["титульный лист", "без таблицы"]}).to_excel(writer, sheet_name="Титул", index=False)
            prices.to_excel(writer, sheet_name="Данные", index=False, startrow=3, startcol=1)
        table, assets, report = self.read(("funds.xlsx", book.getvalue()))
        self.assertEqual(len(table), 72)                                             # 6 лет дневных данных -> 72 месяца
        self.assertAlmostEqual(table["C1"].iloc[-1], prices["Фонд А"].iloc[-1])      # уровень месяца — последнее значение
        self.assertEqual(report["columns"][0]["frequency"], "дневные")
        self.assertIn("skipped", report["files"][0]["tables"][0])                    # титульный лист пропущен с объяснением

    def test_returns_in_percent_with_inflation_column(self):
        returns = self.levels.pct_change().dropna() * 100
        inflation = self.cpi.pct_change().reindex(returns.index) * 100
        text = "Период;Акции, %;Золото, %;Инфляция за месяц, %\n" + "\n".join(
            f"{pd.Timestamp(m):%b %Y};{a:.3f};{b:.3f};{i:.3f}".replace(".", ",") for m, (a, b, _), i in zip(returns.index, returns.values, inflation))
        table, assets, report = self.read(("r.csv", text.encode()))
        self.assertEqual([c["kind"] for c in report["columns"]], ["return_pct"] * 3)
        self.assertEqual(report["inflation"]["mode"], "file")
        self.check(table, "C1", "MCFTR", rtol=1e-3)                                  # из доходностей восстановлен индекс
        cpi = table["CPI"] / table["CPI"].iloc[0]
        self.assertTrue(np.allclose(cpi, self.cpi.reindex(table.index) / self.cpi.reindex(table.index).iloc[0], rtol=1e-3))

    def test_dates_in_a_row_and_year_month_columns(self):
        sideways = self.levels.iloc[:40].T
        table, _, report = self.read(("t.csv", sideways.to_csv(sep=";").encode()))
        self.assertIn("перевёрнута", report["files"][0]["tables"][0]["layout"])
        self.check(table, "C3", "USD")
        fractions = self.levels.pct_change().dropna()
        text = "Год,Месяц,A,B\n" + "\n".join(f"{m[:4]},{int(m[5:])},{a:.5f},{b:.5f}" for m, (a, b, _) in zip(fractions.index, fractions.values))
        table, _, report = self.read(("ym.csv", text.encode()))
        self.assertEqual(report["columns"][0]["kind"], "return_frac")
        self.check(table, "C2", "GOLD", rtol=1e-3)

    def test_rosstat_grid_months_by_years(self):
        """Таблица Росстата: месяцы в строках, годы в столбцах, значения — «в % к предыдущему месяцу»."""
        monthly = (self.cpi.pct_change().dropna() + 1) * 100
        grid = pd.DataFrame({y: {m: monthly.get(f"{y}-{m:02d}") for m in range(1, 13)} for y in range(2013, 2026)})
        grid.index = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]
        book = io.BytesIO()
        with pd.ExcelWriter(book, engine="openpyxl") as writer:
            pd.DataFrame([["Индексы потребительских цен на товары и услуги"]]).to_excel(writer, index=False, header=False)
            grid.to_excel(writer, startrow=3)
        frame, info = reader.extract(reader.read_grids("ipc.xlsx", book.getvalue())[0][1], "ipc")
        self.assertIn("месяцы × годы", info["layout"])
        self.assertEqual(reader.guess_kind(frame.columns[0], frame.iloc[:, 0]), "prev100")
        index = reader.to_monthly_levels(frame.iloc[:, 0], "prev100")
        want = self.cpi.reindex(index.index)
        self.assertTrue(np.allclose(index / index.iloc[0], want / want.iloc[0]))

    def test_html_saved_as_xls_and_user_corrections(self):
        html = "<table><tr><th>Дата</th><th>A</th><th>B</th></tr>" + "".join(
            f"<tr><td>{d:%d.%m.%Y}</td><td>{a:.2f}</td><td>{b:.2f}</td></tr>" for d, (a, b, _) in zip(self.dates, self.levels.values)) + "</table>"
        table, _, report = self.read(("x.xls", html.encode("cp1251")), inflation="none")
        self.assertIn("HTML", report["files"][0]["format"])
        self.assertTrue((table["CPI"] == 1).all())                                   # инфляцию просили не учитывать
        self.check(table, "C1", "MCFTR")

    def test_unsuitable_files_are_rejected_with_reason(self):
        quarterly = "d;a;b\n" + "\n".join(f"{d:%Y-%m-%d};{a};{b}" for d, (a, b, _) in list(zip(self.dates, self.levels.values))[::3])
        for name, raw, reason in (("e.csv", b"", "пустой"), ("g.csv", "привет\nмир".encode(), "не найдена таблица"),
                                  ("q.csv", quarterly.encode(), "квартальные")):
            with self.assertRaises(reader.ReadError) as error:
                self.read((name, raw))
            self.assertIn(reason, str(error.exception))

    def test_number_and_date_formats(self):
        for text, style, value in (("1 234,56", ",", 1234.56), ("1,234.56", ".", 1234.56), ("1.234,56", ",", 1234.56),
                                   ("−0,5%", ",", -0.5), ("(3,2)", ",", -3.2), ("1.2M", ".", 1.2e6), ("12", ".", 12.0)):
            self.assertAlmostEqual(reader.to_number(text, style), value)
        self.assertIsNone(reader.to_number("н/д"))
        for text in ("31.01.2020", "2020-01-31", "2020-01-31 23:59:59", "01.2020", "2020-01", "202001", "20200131",
                     "янв 2020", "Январь 2020", "Jan 31, 2020", "2020M01", "31/01/20"):
            self.assertEqual(reader.to_date(text).strftime("%Y-%m"), "2020-01", text)
        self.assertIsNone(reader.to_date("12345"))
        self.assertEqual(reader.number_style(["5,005", "-0,614", "12,5"]), ",")      # «5,005» двусмысленно, решают соседи
        self.assertEqual(reader.number_style(["1,234.56", "12.5"]), ".")


if __name__ == "__main__":
    unittest.main()
