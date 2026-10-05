"""
Веб-калькулятор портфеля. Запуск на своём компьютере: python app.py  ->  http://127.0.0.1:5000
На Render: gunicorn app:app
"""
import io
import json
import os
import time

import numpy as np
import pandas as pd
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook

import portfolio as pf
import reader

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024          # загружаемые файлы — не больше 20 МБ
app.config["TEMPLATES_AUTO_RELOAD"] = True                    # изменённая страница подхватывается без перезапуска
CRITERIA = {"min_vol": "Минимальная волатильность", "max_sharpe": "Максимальный коэффициент Шарпа",
            "target_return": "Эффективный риск", "target_risk": "Эффективная доходность"}


@app.route("/")
def index():
    # Если в папке static лежит своя копия библиотеки графиков, страница берёт её (работает без интернета)
    local_plotly = os.path.exists(os.path.join(app.static_folder, "plotly.min.js"))
    return render_template("index.html", local_plotly=local_plotly)


def fail(e):
    """Ответ с понятным текстом ошибки. Отдельный код 410 — «сервер забыл загруженную таблицу»:
    получив его, страница сама присылает таблицу заново и повторяет запрос."""
    if isinstance(e, pf.DatasetGone):
        return jsonify({"error": "Загруженная таблица потеряна сервером.", "gone": True}), 410
    return jsonify({"error": str(e)}), 400


@app.errorhandler(413)
def too_large(_):
    return jsonify({"error": "Файл слишком большой: можно загрузить не больше 20 МБ."}), 413


def describe(dataset=None):
    """Список индикаторов и месяцев набора данных — для построения страницы."""
    df, assets = pf.load_data(dataset), pf.assets_of(dataset)
    first = {a["id"]: df[a["id"]].dropna().index[0] for a in assets}
    return {"dataset": dataset, "assets": assets, "first_month": first, "months": list(df.index[1:])}


@app.route("/api/meta")
def meta():
    return jsonify(describe())


@app.route("/api/upload", methods=["POST"])
def upload():
    """Принимает один или несколько файлов Excel/CSV, превращает их в помесячную таблицу и запоминает её.
    В ответе — список индикаторов, отчёт о том, как прочитан файл, и сама таблица (страница хранит её копию)."""
    try:
        files = [(f.filename, f.read()) for f in request.files.getlist("files") if f.filename]
        if not files:
            raise ValueError("Файл не выбран.")
        options = json.loads(request.form.get("options") or "{}")
        table, assets, report = reader.read_files(files, options.get("kinds"), options.get("inflation", "auto"),
                                                  builtin_cpi=pf.builtin_data()["CPI"])
        table = table.round(6)
        dataset = pf.register(table, assets, report)
        columns = {c: [None if pd.isna(v) else float(v) for v in table[c]] for c in table.columns}
        return jsonify({**describe(dataset), "report": report, "kinds": reader.KINDS,
                        "table": {"months": list(table.index), "columns": columns}})
    except Exception as e:  # noqa: BLE001 — показываем пользователю понятное сообщение
        return fail(e)


@app.route("/api/restore", methods=["POST"])
def restore():
    """Страница присылает сохранённую у себя таблицу, если сервер её забыл (например, после перезапуска)."""
    try:
        p = request.get_json(force=True)
        table = pd.DataFrame(p["table"]["columns"], index=p["table"]["months"], dtype=float)
        table = table[[a["id"] for a in p["assets"]] + ["CPI"]]      # столбцы в исходном порядке
        return jsonify({"dataset": pf.register(table, p["assets"], p.get("report"))})
    except Exception as e:  # noqa: BLE001
        return fail(e)


@app.route("/api/template")
def template():
    """Пример файла, который сайт точно поймёт: встроенные данные в виде CSV для Excel."""
    df = pf.builtin_data().rename(columns={**{a["id"]: a["name"] for a in pf.ASSETS}, "CPI": "Индекс потребительских цен (ИПЦ)"})
    df.index.name = "Месяц"
    text = df.to_csv(sep=";", decimal=",", float_format="%.6g")
    return send_file(io.BytesIO(text.encode("utf-8-sig")), as_attachment=True, download_name="primer.csv", mimetype="text/csv")


def read_settings(p):
    """Разбирает и проверяет настройки, присланные страницей."""
    dataset = p.get("dataset") or None                              # номер загруженной таблицы; пусто — встроенные данные
    ids = [a["id"] for a in pf.assets_of(dataset)]
    assets = tuple(a for a in ids if a in p.get("assets", []))
    if len(assets) < 2:
        raise ValueError("Выберите хотя бы два индикатора.")
    criterion = p.get("criterion", "min_vol")
    target = p.get("target")
    if criterion in ("target_return", "target_risk") and target is None:
        raise ValueError("Задайте целевое значение.")
    s = {
        "assets": assets, "criterion": criterion, "target": target,
        "start": p.get("start"), "end": p.get("end"),
        "rf": float(p.get("rf", 0.05)),
        "rent": float(p.get("rent") or 0), "deposit": float(p.get("deposit") or 0),   # доход помимо цены
        "fee": float(p.get("fee") or 0), "tax": float(p.get("tax") or 0),             # издержки
        "dataset": dataset,
    }
    real, _ = pf.prepare(assets, s["start"], s["end"], s["rent"], s["deposit"], dataset)
    if len(real) < 24:
        raise ValueError("Слишком короткий период: нужно не меньше 24 месяцев.")
    return s


def calculate(s):
    """Полный расчёт портфеля по настройкам. Возвращает словарь с результатами.
    По дороге замеряется время каждого этапа и собираются пояснения (trace) для страницы."""
    key = (s["assets"], s["start"], s["end"], s["rent"], s["deposit"], s["dataset"])
    clock = [time.perf_counter()]
    lap = lambda: clock.append(time.perf_counter()) or round((clock[-1] - clock[-2]) * 1000, 1)   # миллисекунды с прошлой отметки

    from_memory = pf.prepare.cache_info().hits
    real, nominal = pf.prepare(*key)
    from_memory = pf.prepare.cache_info().hits > from_memory        # эти данные уже готовили раньше
    ms = {"prepare": lap()}
    mu, cov = pf.annual_stats(real)
    ms["stats"] = lap()
    w, note = pf.optimize(mu, cov, s["criterion"], s["rf"], s["target"])
    ms["optimize"] = lap()
    curve, metrics = pf.curve_and_metrics(real, w, nominal, s["fee"], s["tax"], s["rf"])
    ms["curve"] = lap()
    frontier_from_memory = pf.frontier_for.cache_info().hits
    frontier = pf.frontier_for(*key)
    frontier_from_memory = pf.frontier_for.cache_info().hits > frontier_from_memory
    ms["frontier"] = lap()

    months = pf.load_data(s["dataset"]).index
    start_label = months[months.get_loc(real.index[0]) - 1]          # месяц, когда «вложили рубль»
    labels = [start_label] + list(real.index)
    return {
        "weights": dict(zip(s["assets"], np.round(w, 4).tolist())),
        "metrics": metrics,
        "sharpe": (metrics["mean_return"] - s["rf"]) / metrics["volatility"],
        "note": note,
        "period": [real.index[0], real.index[-1], len(real)],
        "curve": {"months": labels, "values": np.round(curve, 4).tolist()},
        "frontier": frontier,
        "assets": [{"id": a, "risk": float(np.sqrt(cov[i, i])), "ret": float(mu[i])}
                   for i, a in enumerate(s["assets"])],
        # Ориентир: во что за то же время превратился бы рубль, растущий вместе с денежной массой М2 (после инфляции)
        "benchmark": None if s["dataset"] else benchmark_m2(real.index),
        "trace": {
            "ms": ms, "from_memory": from_memory, "frontier_from_memory": frontier_from_memory, "rf": s["rf"],
            "period": pf.explain_period(s["assets"], s["start"], s["end"], real, s["dataset"]),
            "inflation": pf.explain_inflation(real, nominal, s["dataset"]),
            "correlation": np.round(real.corr().values, 3).tolist(),
            "portfolio": pf.explain_portfolio(w, mu, cov),
            "curve": pf.explain_curve(curve, labels, real, nominal, w, s["fee"]),
        },
    }


def benchmark_m2(months):
    """Рост денежной массы М2 сверх инфляции за те же месяцы — линия-ориентир для кривой доходности."""
    m2 = pf.prepare(("M2",), months[0], months[-1])[0]["M2"].reindex(months)
    curve = np.concatenate([[1.0], np.cumprod(1 + m2.values)])
    return {"values": np.round(curve, 4).tolist(), "mean_return": float(m2.mean() * 12)}


@app.route("/api/optimize", methods=["POST"])
def optimize():
    try:
        return jsonify(calculate(read_settings(request.get_json(force=True))))
    except Exception as e:  # noqa: BLE001 — показываем пользователю понятное сообщение
        return fail(e)


@app.route("/api/stability", methods=["POST"])
def stability():
    """Доли портфеля на нескольких сдвинутых периодах."""
    try:
        s = read_settings(request.get_json(force=True))
        windows = pf.stability(s["assets"], s["criterion"], s["rf"], s["target"],
                               s["start"], s["end"], s["rent"], s["deposit"], dataset=s["dataset"])
        return jsonify({"assets": list(s["assets"]), "windows": windows})
    except Exception as e:  # noqa: BLE001
        return fail(e)


@app.route("/api/future", methods=["POST"])
def future():
    """Проверка на «будущем»: доли по ранней части периода, результат — на поздней."""
    try:
        p = request.get_json(force=True)
        s = read_settings(p)
        result = pf.future_check(s["assets"], s["criterion"], s["rf"], s["target"], s["start"], s["end"],
                                 s["rent"], s["deposit"], p.get("split"), s["dataset"])
        return jsonify({"assets": list(s["assets"]), **result})
    except Exception as e:  # noqa: BLE001
        return fail(e)


@app.route("/api/export", methods=["POST"])
def export():
    """Тот же расчёт, но ответ — файл Excel с четырьмя листами."""
    try:
        s = read_settings(request.get_json(force=True))
        r = calculate(s)
    except Exception as e:  # noqa: BLE001
        return fail(e)
    m = r["metrics"]
    NAMES = {a["id"]: a["name"] for a in pf.assets_of(s["dataset"])}
    wb = Workbook()

    ws = wb.active
    ws.title = "Портфель"
    rows = [
        ("Критерий", CRITERIA[s["criterion"]]),
        ("Период", f"{r['period'][0]} — {r['period'][1]} ({r['period'][2]} мес.)"),
        ("Безрисковая ставка", s["rf"]),
        ("Целевое значение", s["target"] if s["target"] is not None else "—"),
        ("Аренда, в год", s["rent"]), ("Ставка по валютному вкладу, в год", s["deposit"]),
        ("Комиссия с оборота", s["fee"]), ("Налог на доход", s["tax"]),
        (),
        ("Ожидаемый доход, в год", m["mean_return"]), ("Волатильность", m["volatility"]),
        ("Риск потерь", m["downside"]), ("Макс. просадка", m["max_drawdown"]),
        ("Макс. период восстановления, мес.", m["max_recovery_months"]),
        ("Коэффициент Шарпа", r["sharpe"]), ("Коэффициент Сортино", m["sortino"]),
        ("1 рубль превратился в", m["final_value"]), ("То же после налога", m["final_after_tax"]),
        (),
        ("Индикатор", "Доля"),
    ] + [(NAMES[a], w) for a, w in sorted(r["weights"].items(), key=lambda x: -x[1]) if w > 0]
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 36

    ws = wb.create_sheet("Кривая доходности")
    ws.append(("Месяц", "Стоимость 1 рубля"))
    for row in zip(r["curve"]["months"], r["curve"]["values"]):
        ws.append(row)

    ws = wb.create_sheet("Эффективная граница")
    ws.append(("Риск", "Доходность"))
    for row in r["frontier"]:
        ws.append(tuple(row))

    ws = wb.create_sheet("Индикаторы")
    ws.append(("Индикатор", "Доходность, в год", "Риск, в год"))
    for a in r["assets"]:
        ws.append((NAMES[a["id"]], a["ret"], a["risk"]))
    ws.column_dimensions["A"].width = 44

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="portfolio.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    # На своём компьютере сайт виден только с него самого: http://127.0.0.1:5000
    # use_reloader: если файлы программы изменились, сервер перезапускается сам
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", 5000)), use_reloader=True)
